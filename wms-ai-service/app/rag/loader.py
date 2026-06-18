import logging
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

logger = logging.getLogger(__name__)

# 文件大小限制：50MB
MAX_FILE_SIZE = 50 * 1024 * 1024


def load_document(path: Path) -> str:
    if not path.exists():
        raise FileNotFoundError(f"file not found: {path}")

    # 检查文件大小
    file_size = path.stat().st_size
    if file_size > MAX_FILE_SIZE:
        raise ValueError(f"文件大小超过限制: {file_size / 1024 / 1024:.1f}MB > {MAX_FILE_SIZE / 1024 / 1024}MB")

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return load_pdf(path)
        if suffix in {".docx", ".doc"}:
            return load_docx(path)
        if suffix in {".xlsx", ".xls"}:
            return load_xlsx(path)
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:
        logger.error("Failed to load document %s: %s", path, exc)
        raise ValueError(f"文件解析失败: {exc}") from exc


def load_pdf(path: Path) -> str:
    try:
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ValueError(f"PDF解析失败: {exc}") from exc


def load_docx(path: Path) -> str:
    try:
        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    except Exception as exc:
        raise ValueError(f"DOCX解析失败: {exc}") from exc


def load_xlsx(path: Path) -> str:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell is not None]
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)
    except Exception as exc:
        raise ValueError(f"XLSX解析失败: {exc}") from exc
